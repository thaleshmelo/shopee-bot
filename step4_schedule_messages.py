import os
import pandas as pd
from datetime import datetime
from pathlib import Path

TOTAL_MESSAGES = 15
TIME_BLOCKS = [("09:00", "14:00", 5), ("14:00", "20:00", 5), ("20:00", "23:30", 5)]

MESSAGES_FILE = os.getenv("SHOPEE_MESSAGES_FILE", "").strip()
PICKS_FILE = os.getenv("SHOPEE_PICKS_FILE", "picks_refinados.csv").strip()

BLOCKLIST_TERMS = [
    "vibrador", "bullet", "sexy", "lingerie", "erótico", "erotico", "sexshop",
    "patê", "pate", "pimenta", "alimento", "comida", "whey", "creatina", "psyllium",
    "suplemento", "vitamina", "capsula", "cápsula",
    "xbox 360", "ps3", "playstation 3", "wii",
]

def parse_time(t: str) -> datetime:
    today = datetime.now().date()
    h, m = map(int, t.split(":"))
    return datetime.combine(today, datetime.min.time()).replace(hour=h, minute=m)

def distribute_times():
    times = []
    for start, end, qty in TIME_BLOCKS:
        start_dt = parse_time(start)
        end_dt = parse_time(end)
        step = (end_dt - start_dt) / qty
        for i in range(qty):
            times.append(start_dt + step * i)
    return times[:TOTAL_MESSAGES]

def normalize_cols(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df.columns = [str(c).strip().replace(" ", "_") for c in df.columns]
    return df

def looks_blocked(text: str) -> bool:
    t = (text or "").lower()
    return any(term in t for term in BLOCKLIST_TERMS)

def read_input_dataframe() -> pd.DataFrame:
    if MESSAGES_FILE:
        if not Path(MESSAGES_FILE).exists():
            raise SystemExit(f"SHOPEE_MESSAGES_FILE não encontrado: {MESSAGES_FILE}")
        if MESSAGES_FILE.lower().endswith(".csv"):
            return normalize_cols(pd.read_csv(MESSAGES_FILE, low_memory=False))
        if MESSAGES_FILE.lower().endswith(".xlsx"):
            return normalize_cols(pd.read_excel(MESSAGES_FILE))
        raise SystemExit("SHOPEE_MESSAGES_FILE deve ser .csv ou .xlsx")

    if not Path(PICKS_FILE).exists():
        raise SystemExit("Defina SHOPEE_MESSAGES_FILE ou garanta picks_refinados.csv (Step2).")
    return normalize_cols(pd.read_csv(PICKS_FILE, low_memory=False))

def build_messages_from_picks(df: pd.DataFrame) -> pd.DataFrame:
    title_col = "titulo" if "titulo" in df.columns else "title"
    price_col = "preco" if "preco" in df.columns else "sale_price"
    disc_col  = "desconto_pct" if "desconto_pct" in df.columns else "discount_percentage"

    link_col = None
    for c in ["product_short_link", "product_link", "link"]:
        if c in df.columns:
            link_col = c
            break
    if not link_col:
        for c in df.columns:
            if "link" in c.lower():
                link_col = c
                break
    if not link_col:
        raise SystemExit("Nenhuma coluna de link encontrada no arquivo de picks.")

    out = []
    for _, r in df.iterrows():
        title = str(r.get(title_col, "")).strip()
        link = str(r.get(link_col, "")).strip()

        try:
            price = float(str(r.get(price_col, "0")).replace(",", "."))
        except Exception:
            price = 0.0
        try:
            disc = int(float(str(r.get(disc_col, "0")).replace(",", ".")))
        except Exception:
            disc = 0

        msg = (
            f"🔥 {title}\n"
            f"💰 R$ {price:.2f} (-{disc}%)\n"
            f"🔗 {link}"
        )
        out.append({"mensagem": msg, "titulo": title, "link": link})
    return pd.DataFrame(out)

def save_xlsx_pretty(df: pd.DataFrame, filename: str):
    # Requer openpyxl
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    df.to_excel(filename, index=False)
    wb = load_workbook(filename)
    ws = wb.active

    # Header
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center")

    # Larguras
    widths = {"hora": 9, "mensagem": 70}
    for col_idx, col_name in enumerate(df.columns, start=1):
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = widths.get(col_name, 18)

    # Alinhamentos e wrap
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.column_letter == "A":  # hora
                cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=False)
            else:  # mensagem
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    wb.save(filename)

def main():
    df = read_input_dataframe()

    if "mensagem" in df.columns:
        msgs = df.copy()
        msgs["mensagem"] = msgs["mensagem"].astype(str)
        if "titulo" not in msgs.columns:
            msgs["titulo"] = ""
    else:
        msgs = build_messages_from_picks(df)

    msgs = msgs[~msgs.apply(lambda r: looks_blocked(str(r.get("titulo",""))) or looks_blocked(str(r.get("mensagem",""))), axis=1)].copy()
    msgs = msgs.head(TOTAL_MESSAGES).copy()

    schedule = distribute_times()

    rows = []
    print("\n===== AGENDA DE ENVIOS (15/dia) =====\n")
    for i, (t, (_, r)) in enumerate(zip(schedule, msgs.iterrows()), start=1):
        time_str = t.strftime("%H:%M")
        msg = str(r["mensagem"]).strip()
        print(f"[{time_str}] Mensagem {i}\n{msg}\n" + "-" * 45)
        rows.append({"hora": time_str, "mensagem": msg})

    out = pd.DataFrame(rows)
    base = f"agenda_envios_{datetime.now().strftime('%Y-%m-%d')}"

    csv_name = f"{base}.csv"
    out.to_csv(csv_name, index=False, encoding="utf-8-sig")

    # XLSX “bonito”
    xlsx_name = f"{base}.xlsx"
    save_xlsx_pretty(out, xlsx_name)

    print(f"\n📁 Salvo: {csv_name}")
    print(f"📁 Salvo: {xlsx_name} (formatado)")

if __name__ == "__main__":
    main()
